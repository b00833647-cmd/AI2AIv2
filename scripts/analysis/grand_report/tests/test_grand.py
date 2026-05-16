"""Deterministic, offline tests for the grand four-experiment report.

Covers the core mandate (zero prior-report imports), the X1–X4
framework, the modern stats, the quad runner at all four levels, every
group, and the assembled APA-7 manuscript (structure + auto-ToC +
honesty charter).
"""
from __future__ import annotations

import ast
import os
import pathlib
import re
import zipfile

import pytest
from docx import Document

from scripts.analysis.grand_report import db
from scripts.analysis.grand_report import framework as FW
from scripts.analysis.grand_report import metrics as M
from scripts.analysis.grand_report import report as RP
from scripts.analysis.grand_report import stats as S

PKG = pathlib.Path(__file__).resolve().parents[1]
FORBIDDEN = ("human_pilot", "process_report", "confirmatory",
             "apa_report", "analysis.core")
pytestmark = pytest.mark.skipif(not os.path.exists(db.SNAPSHOT),
                                reason="snapshot absent")


def test_zero_prior_report_imports():
    for f in list(PKG.glob("*.py")) + list(PKG.glob("tests/*.py")):
        tree = ast.parse(f.read_text(), filename=str(f))
        for node in ast.walk(tree):
            mods = []
            if isinstance(node, ast.Import):
                mods = [a.name for a in node.names]
            elif isinstance(node, ast.ImportFrom):
                mods = [node.module or ""]
            for m in mods:
                for bad in FORBIDDEN:
                    assert bad not in m, f"{f.name} imports {m!r}"


def test_four_experiment_framework():
    con = db.connect()
    try:
        c = db.completers(con)
        assert len(c) == 40
        assert (c.groupby("exp").size().reindex(
            ["X1", "X2", "X3", "X4"]) == 10).all()
        assert db.verify_reservations(con)["seller_floor"] == 22500
        w = db.survey_wide(con)
        assert all(k in w.columns for k in db.SURVEY_KEYS)
    finally:
        con.close()


def test_modern_stats():
    assert S.cliffs_delta([9, 9, 9], [1, 1, 1]) == 1.0
    q = S.multipletests([0.2, 0.01, 0.04], "fdr_bh")
    assert all(not (x != x) for x in q)
    mw = S.mann_whitney([1, 2, 1, 2], [6, 7, 6, 7])
    assert "band" in mw and mw["p"] < 0.05
    srh = S.scheirer_ray_hare(
        __import__("pandas").DataFrame(
            {"v": range(16), "a": ["x", "y"] * 8, "b": ["p", "q"] * 8}),
        "v", "a", "b")
    assert set(srh) == {"A", "B", "AB", "N"}
    assert "bf10" in S.bayes_ttest([1, 2, 3, 4], [2, 3, 4, 5])
    assert "p_tost" in S.tost([1, 2, 3, 4], [1, 2, 3, 4])


def test_quad_runner_levels():
    con = db.connect()
    try:
        pf = M.g04(con)["participant"]
        q = FW.quad_continuous(pf, "satisfaction")
        assert {"l1", "l2_factorial", "l2_pairwise", "l3"} <= set(q)
        assert len(q["l1"]) == 4  # X1–X4 one-by-one
        assert len(q["l2_pairwise"]) == 6  # structured pairwise
        mq = FW.multi_quad(pf, db.SURVEY_KEYS)
        assert len(mq["l1"]) == 9 and {"l1", "l2", "l3"} <= set(mq)
        qb = FW.quad_binary(M.g03b_outcomes(con)["participant"], "agreed")
        assert len(qb["l1"]) == 4
    finally:
        con.close()


def test_all_groups_run():
    con = db.connect()
    try:
        assert len(M.GROUPS) >= 45
        for fn in M.GROUPS:
            r = fn(con)
            assert r["id"] and r["title"] and r["chapter"] in M.CHAPTERS
            pf = r["participant"]
            if pf is not None and (r["value_cols"] or r["binary_cols"]):
                assert {"exp", "mode", "role", "opponent"} <= set(
                    pf.columns)
        assert len(M.g_econ(con)["participant"]) == 27
        assert len(M.g04(con)["value_cols"]) == 9
        assert len(M.g10(con)["tables"]["Verbatim (anonymised)"]) == 40
    finally:
        con.close()


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("grand") / "g.docx"
    res = RP.build(snapshot=db.SNAPSHOT, out=out)
    doc = Document(str(out))
    txt = "\n".join(p.text for p in doc.paragraphs)
    with zipfile.ZipFile(str(out)) as z:
        media = [m for m in z.namelist()
                 if m.startswith("word/media/")]
        docxml = z.read("word/document.xml").decode("utf8", "ignore")
        hdr = "".join(
            z.read(n).decode("utf8", "ignore") for n in z.namelist()
            if re.match(r"word/header\d*\.xml", n))
    return dict(res=res, doc=doc, txt=txt, media=media,
                docxml=docxml, hdr=hdr)


def test_manuscript_structure(built):
    assert built["res"]["bytes"] > 1_000_000
    assert len(built["doc"].tables) >= 120
    assert len(built["media"]) >= 40
    h1 = [p.text for p in built["doc"].paragraphs
          if p.style.name == "Heading 1"]
    for part in ["Part I — Foundations",
                 "Part II — The Four Experiments (One by One)",
                 "Part III — The Forty-Five Analyses",
                 "Part IV — Cross-Experiment Synthesis, Robustness & "
                 "Confirmatory Roadmap", "References"]:
        assert part in h1
    assert any(x.startswith("Appendix A") for x in h1)
    assert any(x.startswith("Appendix B") for x in h1)


def test_auto_toc_and_page_numbers(built):
    assert 'TOC \\o' in built["docxml"]
    assert 'TOC \\h \\z \\c "Table"' in built["docxml"]
    assert 'TOC \\h \\z \\c "Figure"' in built["docxml"]
    assert "SEQ Table" in built["docxml"]
    assert "SEQ Figure" in built["docxml"]
    assert "PAGE" in built["hdr"]


def test_four_levels_present(built):
    t = built["txt"]
    assert t.count("L1 Distinct") >= 20
    assert t.count("L2 2×2") >= 20
    assert t.count("L3 Overall") >= 20
    assert t.count("L4 review") >= 40


def test_honesty_charter(built):
    t = built["txt"]
    assert "Amir Sepehri" in t and "b00833647@essec.edu" in t
    assert "Exploratory" in RP.TITLE
    assert "never a pre-registered confirmatory result" in t
    assert "randomised, not controlled" in t
    assert "utility and judge-score tables are empty" in t
    assert not re.search(r"\[-0\.\d", t)  # APA CI formatting holds
