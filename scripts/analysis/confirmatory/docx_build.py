"""Thin python-docx helper for the confirmatory report (python-docx only)."""
from __future__ import annotations

from docx import Document


def new_doc():
    return Document()


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def para(doc, text, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if italic:
        run.italic = True
    return p


def table(doc, df, caption):
    cap = doc.add_paragraph()
    cap.add_run(caption).bold = True
    cols = [str(c) for c in df.columns]
    t = doc.add_table(rows=1 + len(df), cols=max(1, len(cols)))
    try:
        t.style = "Light Grid Accent 1"
    except Exception:
        pass
    for j, name in enumerate(cols):
        t.rows[0].cells[j].text = name
    for i, row in enumerate(df.itertuples(index=False, name=None), start=1):
        for j, value in enumerate(row):
            t.rows[i].cells[j].text = str(value)
