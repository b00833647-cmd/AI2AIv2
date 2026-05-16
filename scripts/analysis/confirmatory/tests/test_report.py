import os
import sqlite3
import pandas as pd
import pytest
from scripts.analysis.confirmatory.report import build_report, DRY_RUN_STAMP

PILOT = "data/ai2ai-human-pilot-2026-05-15.db"


def _clean_db(tmp_path):
    p = tmp_path / "clean.db"
    c = sqlite3.connect(p)
    c.executescript("""
      CREATE TABLE study_participants (
        id TEXT PRIMARY KEY, role TEXT, completion_code TEXT, experiment_mode TEXT,
        session_id TEXT, condition_mode TEXT, condition_role TEXT, opponent_block TEXT,
        assignment_seed TEXT, assignment_block_index INTEGER, replicate_id INTEGER,
        manipulation_check_pass INTEGER, excl_attention INTEGER,
        excl_manipulation INTEGER, excl_speeding INTEGER, excl_comprehension INTEGER,
        excl_noncompletion INTEGER);
      CREATE TABLE participant_responses (
        participant_id TEXT, screen TEXT, key TEXT, value_int INTEGER, value_text TEXT);
      CREATE TABLE sessions (id TEXT PRIMARY KEY, outcome_type TEXT, outcome_terms_json TEXT);
    """)
    import numpy as np
    rng = np.random.default_rng(0)
    n = 0
    for m in ("delegated", "direct"):
        for r in ("buyer", "seller"):
            for o in ("easygoing", "moderate", "tough"):
                for _ in range(8):
                    pid = f"P{n}"; n += 1
                    c.execute("INSERT INTO study_participants VALUES "
                              "(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                              (pid, r, "CC", "agent", f"s{pid}", m, r, o,
                               "SEED", 0, 0, 1, 0, 0, 0, 0, 0))
                    c.execute("INSERT INTO sessions VALUES (?,?,?)",
                              (f"s{pid}", "agreed", '{"price": 23000}'))
                    sat = int(np.clip(round(rng.normal(5.4 if m == "delegated" else 4.2, 1)), 1, 7))
                    for k in ("satisfaction","would_use_again","agent_represented",
                              "control","emot_pleasant","emot_anxious",
                              "effort_invested","engage_engaged","outfair_share"):
                        c.execute("INSERT INTO participant_responses VALUES (?,?,?,?,NULL)",
                                  (pid, "post_survey", k, sat if k=="satisfaction" else 4))
    c.commit(); c.close()
    return str(p)


def test_clean_db_produces_unstamped_report(tmp_path):
    out = tmp_path / "conf.docx"
    res = build_report(_clean_db(tmp_path), str(out))
    assert res["clean"] is True
    assert res["dry_run"] is False
    assert out.exists()
    from docx import Document
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert DRY_RUN_STAMP not in text
    assert "satisfaction" in text.lower()


@pytest.mark.skipif(not os.path.exists(PILOT), reason="frozen pilot DB absent")
def test_pilot_is_dry_run_and_stamped(tmp_path):
    out = tmp_path / "pilot_conf.docx"
    res = build_report(PILOT, str(out))
    assert res["clean"] is False
    assert res["dry_run"] is True
    assert out.exists()
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert DRY_RUN_STAMP in text
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any(DRY_RUN_STAMP in h for h in heads)
