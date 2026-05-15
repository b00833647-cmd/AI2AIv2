"""Clean-room python-docx assembler for the process report. python-docx only."""
from __future__ import annotations

from docx import Document
from docx.shared import Inches


def new_doc():
    """A fresh, empty Word document."""
    return Document()


def h1(doc, text):
    """Level-1 heading."""
    return doc.add_heading(text, level=1)


def h2(doc, text):
    """Level-2 heading."""
    return doc.add_heading(text, level=2)


def para(doc, text, italic=False):
    """A paragraph; italicize its run when requested. Returns the paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    if italic:
        run.italic = True
    return p


def figure(doc, path, caption, note):
    """Embed an image (6in wide) with a bold caption and an italic note."""
    doc.add_picture(str(path), width=Inches(6.0))
    cap = doc.add_paragraph()
    cap.add_run(caption).bold = True
    para(doc, note, italic=True)


def table(doc, df, caption, note):
    """A captioned table from a DataFrame, then an italic note.

    Header row = stringified columns; one row per record (cells stringified).
    An empty DataFrame yields a header-only table without crashing.
    """
    cap = doc.add_paragraph()
    cap.add_run(caption).bold = True
    cols = [str(c) for c in df.columns]
    t = doc.add_table(rows=1 + len(df), cols=len(cols))
    try:
        t.style = "Light Grid Accent 1"
    except Exception:
        pass
    for j, name in enumerate(cols):
        t.rows[0].cells[j].text = name
    for i, row in enumerate(df.itertuples(index=False, name=None), start=1):
        for j, value in enumerate(row):
            t.rows[i].cells[j].text = str(value)
    para(doc, note, italic=True)
