"""python-docx block helpers: heading, figure+caption+interpretation, table."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


def new_doc() -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    return doc


def h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(16)
    p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)


def h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1B, 0x3F, 0x8B)
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)


def para(doc, text, *, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = italic
    if italic:
        r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def figure(doc, png_path: Path, caption: str, interpretation: str):
    pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(str(png_path), width=Inches(6.0))
    cap = doc.add_paragraph()
    cr = cap.add_run(caption); cr.italic = True; cr.font.size = Pt(9)
    para(doc, interpretation)


def table(doc, df: pd.DataFrame, caption: str, interpretation: str):
    cap = doc.add_paragraph()
    cr = cap.add_run(caption); cr.bold = True; cr.font.size = Pt(10)
    t = doc.add_table(rows=1 + len(df), cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for j, col in enumerate(df.columns):
        c = t.rows[0].cells[j]; c.text = ""
        rr = c.paragraphs[0].add_run(str(col)); rr.bold = True; rr.font.size = Pt(9)
    for i, (_, row) in enumerate(df.iterrows()):
        for j, col in enumerate(df.columns):
            cell = t.rows[i + 1].cells[j]; cell.text = ""
            rr = cell.paragraphs[0].add_run("" if pd.isna(row[col]) else str(row[col]))
            rr.font.size = Pt(9)
    doc.add_paragraph()
    para(doc, interpretation)
